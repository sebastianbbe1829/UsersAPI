import io
from datetime import UTC, date, datetime, time
from zoneinfo import ZoneInfo

from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.orm import Session, joinedload

from ..models import InventoryDB, InventoryMovementDB, ProductDB
from .inventory_service import _with_calculated_values

COLOMBIA_TZ = ZoneInfo("America/Bogota")
ORIGIN_LABELS = {
    "PURCHASE": "Compra",
    "SALE": "Venta",
    "MANUAL_ADJUSTMENT": "Ajuste de inventario",
    "SALES_RETURN": "Devolución de venta",
    "PURCHASE_RETURN": "Devolución de compra",
    "REVERSAL": "Reversión",
}


def _response(wb, filename):
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _sheet(ws, title, headers):
    ws.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=len(headers),
    )
    ws.cell(1, 1, title).font = Font(bold=True, size=16, color="FFFFFF")
    ws.cell(1, 1).fill = PatternFill("solid", fgColor="17365D")
    ws.cell(1, 1).alignment = Alignment(horizontal="center")
    for col, header in enumerate(headers, 1):
        cell = ws.cell(3, col, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A4"
    ws.sheet_view.showGridLines = False


def export_inventory_excel(
    db: Session,
    tenant_id: int,
    search: str | None = None,
    inventory_type_id: int | None = None,
):
    query = (
        db.query(InventoryDB)
        .options(joinedload(InventoryDB.product).joinedload(ProductDB.inventory_type))
        .join(
            ProductDB,
            (ProductDB.id == InventoryDB.product_id) & (ProductDB.tenant_id == tenant_id),
        )
        .filter(InventoryDB.tenant_id == tenant_id)
    )
    if search:
        term = f"%{search.strip()}%"
        query = query.filter((ProductDB.code.ilike(term)) | (ProductDB.name.ilike(term)))
    if inventory_type_id is not None:
        query = query.filter(ProductDB.inventory_type_id == inventory_type_id)
    rows = query.order_by(ProductDB.name, ProductDB.id).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventario"
    headers = [
        "Código",
        "Producto",
        "Tipo",
        "Existencia",
        "Precio compra",
        "Ganancia",
        "Precio venta",
        "Total inventario",
    ]
    _sheet(ws, "REPORTE DE INVENTARIO", headers)
    for row, inventory in enumerate(rows, 4):
        item = _with_calculated_values(inventory)
        product = inventory.product
        values = [
            product.code if product else "",
            product.name if product else "",
            product.inventory_type.name if product and product.inventory_type else "",
            float(item.quantity or 0),
            float(item.purchase_price) if item.purchase_price is not None else None,
            float(item.profit_percentage or 0),
            float(item.sale_price) if item.sale_price is not None else None,
            float(item.total_inventory or 0),
        ]
        for col, value in enumerate(values, 1):
            ws.cell(row, col, value)
        ws.cell(row, 6).number_format = "0.00%"
        for col in (5, 7, 8):
            ws.cell(row, col).number_format = "#,##0.00"
    if rows:
        ws.auto_filter.ref = f"A3:H{len(rows) + 3}"
    for col, width in enumerate([18, 36, 28, 14, 18, 14, 18, 20], 1):
        ws.column_dimensions[chr(64 + col)].width = width
    return _response(wb, "inventario.xlsx")


def export_movements_excel(
    db: Session,
    tenant_id: int,
    product_id: int | None = None,
    from_date: date | None = None,
    to_date: date | None = None,
):
    if from_date and to_date and from_date > to_date:
        raise ValueError("from_date cannot be greater than to_date")
    start = (
        datetime.combine(from_date, time.min, tzinfo=COLOMBIA_TZ)
        .astimezone(UTC)
        .replace(tzinfo=None)
        if from_date
        else None
    )
    end = (
        datetime.combine(to_date, time.max, tzinfo=COLOMBIA_TZ).astimezone(UTC).replace(tzinfo=None)
        if to_date
        else None
    )
    query = (
        db.query(InventoryMovementDB)
        .options(joinedload(InventoryMovementDB.product).joinedload(ProductDB.inventory_type))
        .filter(InventoryMovementDB.tenant_id == tenant_id)
    )
    if product_id is not None:
        query = query.filter(InventoryMovementDB.product_id == product_id)
    if start is not None:
        query = query.filter(InventoryMovementDB.created_at >= start)
    if end is not None:
        query = query.filter(InventoryMovementDB.created_at <= end)
    rows = query.order_by(InventoryMovementDB.created_at, InventoryMovementDB.id).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Kardex"
    headers = [
        "Fecha",
        "Movimiento",
        "Operación",
        "Producto",
        "Código",
        "Tipo",
        "Cantidad",
        "Precio compra",
        "Antes",
        "Después",
        "ID movimiento",
        "ID movimiento original",
        "Notas",
    ]
    _sheet(ws, "REPORTE DE KARDEX", headers)
    for row, movement in enumerate(rows, 4):
        product = movement.product
        values = [
            movement.created_at,
            "Entrada" if movement.movement_type == "ENTRY" else "Salida",
            ORIGIN_LABELS.get(movement.origin_type, movement.origin_type),
            product.name if product else "",
            product.code if product else "",
            product.inventory_type.name if product and product.inventory_type else "",
            float(movement.quantity),
            (
                float(movement.unit_purchase_price)
                if movement.unit_purchase_price is not None
                else None
            ),
            float(movement.balance_before),
            float(movement.balance_after),
            str(movement.id),
            str(movement.reversal_of_id) if movement.reversal_of_id else "",
            movement.notes or "",
        ]
        for col, value in enumerate(values, 1):
            ws.cell(row, col, value)
        ws.cell(row, 1).number_format = "dd/mm/yyyy hh:mm:ss"
        for col in (7, 9, 10):
            ws.cell(row, col).number_format = "#,##0.000"
        ws.cell(row, 8).number_format = "#,##0.00"
    if rows:
        ws.auto_filter.ref = f"A3:M{len(rows) + 3}"
    for col, width in enumerate(
        [20, 14, 24, 34, 18, 26, 14, 18, 14, 14, 38, 38, 45],
        1,
    ):
        ws.column_dimensions[chr(64 + col)].width = width
    return _response(wb, "kardex.xlsx")
