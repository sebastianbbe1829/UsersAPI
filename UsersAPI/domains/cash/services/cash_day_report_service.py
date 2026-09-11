from collections import defaultdict
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import select
from sqlalchemy.orm import Session

from UsersAPI.domains.portfolio.models import PaymentDB
from UsersAPI.domains.sales.models import SaleDB, SalePaymentDB

from ..models import BranchDB, CashBoxDB, CashDayDB, CashMovementDB, CashRegisterDB

ZERO = Decimal("0.00")
CASH = {"CASH", "EFECTIVO"}
PREFERRED_METHODS = ["Efectivo", "Transferencias", "PSE", "Tarjeta", "Crédito"]
COLOMBIA_TZ = ZoneInfo("America/Bogota")


def _money(value):
    return Decimal(str(value or 0)).quantize(Decimal("0.01"))


def _method(value):
    method = (value or "OTRO").strip().upper()
    if method in CASH:
        return "Efectivo"
    if method in {"TRANSFER", "TRANSFERENCIA", "BANK_TRANSFER"}:
        return "Transferencias"
    if method in {"CARD", "TARJETA", "CREDIT_CARD", "DEBIT_CARD"}:
        return "Tarjeta"
    if method in {"CREDIT", "CREDITO", "CRÉDITO"}:
        return "Crédito"
    if method == "PSE":
        return "PSE"
    return method.title()


def _ordered_methods(sales_day, payments_day):
    available = set(sales_day) | set(payments_day)
    ordered = [method for method in PREFERRED_METHODS if method in available]
    extras = sorted(available - set(PREFERRED_METHODS))
    return ordered + extras


def build_day_report(
    db: Session,
    tenant_id: int,
    day_id: int,
    *,
    tenant_name: str | None = None,
    generated_by: str | None = None,
    generated_at: datetime | None = None,
) -> dict:
    day = db.scalar(select(CashDayDB).where(CashDayDB.tenant_id == tenant_id, CashDayDB.id == day_id))
    if day is None:
        raise ValueError("Día operativo no encontrado.")

    registers = db.scalars(
        select(CashRegisterDB).where(
            CashRegisterDB.tenant_id == tenant_id,
            CashRegisterDB.cash_day_id == day.id,
        ).order_by(CashRegisterDB.branch_id, CashRegisterDB.id)
    ).all()
    branches = {b.id: b.name for b in db.scalars(select(BranchDB).where(BranchDB.tenant_id == tenant_id)).all()}
    boxes = {b.id: b.name for b in db.scalars(select(CashBoxDB).where(CashBoxDB.tenant_id == tenant_id)).all()}
    movements = db.scalars(
        select(CashMovementDB).where(
            CashMovementDB.tenant_id == tenant_id,
            CashMovementDB.business_date == day.business_date,
        ).order_by(CashMovementDB.cash_register_id, CashMovementDB.id)
    ).all()

    sales_by_box = defaultdict(lambda: defaultdict(lambda: ZERO))
    payments_by_box = defaultdict(lambda: defaultdict(lambda: ZERO))
    cash_delta_by_register = defaultdict(lambda: ZERO)
    cash_sales = ZERO
    cash_payments = ZERO
    cash_sales_reversals = ZERO
    cash_payment_reversals = ZERO
    for movement in movements:
        amount = _money(movement.amount)
        signed = amount if movement.movement_type == "INCOME" else -amount
        method = _method(movement.payment_method)
        if method == "Efectivo":
            cash_delta_by_register[movement.cash_register_id] += signed
        if movement.origin_type == "SALE":
            sales_by_box[movement.cash_register_id][method] += signed
            if method == "Efectivo":
                if movement.movement_type == "INCOME":
                    cash_sales += amount
                else:
                    cash_sales_reversals += amount
        elif movement.origin_type in {"PORTFOLIO_PAYMENT", "PAYMENT_REVERSAL"}:
            payments_by_box[movement.cash_register_id][method] += signed
            if method == "Efectivo":
                if movement.origin_type == "PORTFOLIO_PAYMENT":
                    cash_payments += amount
                else:
                    cash_payment_reversals += amount

    sale_rows = db.execute(
        select(SalePaymentDB.payment_method, SalePaymentDB.amount).join(
            SaleDB, SaleDB.id == SalePaymentDB.sale_id
        ).where(
            SalePaymentDB.tenant_id == tenant_id,
            SaleDB.tenant_id == tenant_id,
            SaleDB.business_date == day.business_date,
            SaleDB.status != "CANCELLED",
        )
    ).all()
    sales_day = defaultdict(lambda: ZERO)
    for method, amount in sale_rows:
        sales_day[_method(method)] += _money(amount)

    payment_rows = db.execute(
        select(PaymentDB.payment_method, PaymentDB.amount, PaymentDB.status).where(
            PaymentDB.tenant_id == tenant_id,
            PaymentDB.payment_date == day.business_date,
        )
    ).all()
    payments_day = defaultdict(lambda: ZERO)
    for method, amount, status in payment_rows:
        payments_day[_method(method)] += _money(amount) if status == "APLICADO" else -_money(amount)

    register_rows = []
    total_expected = ZERO
    total_counted = ZERO
    total_difference = ZERO
    closed_ok = []
    closed_mismatch = []
    pending = []
    for register in registers:
        base_amount = _money(register.opening_amount)
        expected_calculated = base_amount + cash_delta_by_register[register.id]
        expected = expected_calculated if register.status == "OPEN" or register.expected_cash is None else _money(register.expected_cash)
        counted = _money(register.counted_cash) if register.counted_cash is not None else None
        difference = _money(register.difference) if register.difference is not None else None
        total_expected += expected
        if counted is not None:
            total_counted += counted
            total_difference += difference or ZERO
        row = {
            "register_id": register.id,
            "branch": branches.get(register.branch_id, "—"),
            "box": boxes.get(register.cash_box_id, f"Caja #{register.id}"),
            "status": register.status,
            "base_amount": base_amount,
            "opening_amount": base_amount,
            "opened_at": register.opened_at,
            "closed_at": register.closed_at,
            "expected": expected,
            "counted": counted,
            "difference": difference,
            "sales": dict(sales_by_box[register.id]),
            "payments": dict(payments_by_box[register.id]),
        }
        register_rows.append(row)
        if register.status == "CLOSED" and difference is not None:
            (closed_ok if difference == ZERO else closed_mismatch).append(row)
        else:
            pending.append(row)

    methods = _ordered_methods(sales_day, payments_day)
    generated_at = generated_at or datetime.now(COLOMBIA_TZ)
    return {
        "day": day,
        "registers": register_rows,
        "methods": methods,
        "sales_day": dict(sales_day),
        "payments_day": dict(payments_day),
        "closed_ok": closed_ok,
        "closed_mismatch": closed_mismatch,
        "pending": pending,
        "total_expected": total_expected,
        "total_counted": total_counted,
        "total_difference": total_difference,
        "cash_sales": cash_sales,
        "cash_payments": cash_payments,
        "cash_sales_reversals": cash_sales_reversals,
        "cash_payment_reversals": cash_payment_reversals,
        "tenant_name": tenant_name or "—",
        "generated_by": generated_by or "—",
        "generated_at": generated_at,
    }


def _filename(report, extension):
    return f"cierre_caja_{report['day'].business_date.isoformat()}.{extension}"


def excel_report(report) -> tuple[bytes, str]:
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws.append(["RESUMEN CIERRE DE CAJA", None])
    ws.append(["Tenant", report["tenant_name"]])
    ws.append(["Fecha operativa", report["day"].business_date.isoformat()])
    ws.append(["Estado", report["day"].status])
    ws.append(["Generado por", report["generated_by"]])
    ws.append(["Fecha/hora generación", report["generated_at"].strftime("%d/%m/%Y %H:%M:%S")])
    ws.append(["Apertura del día", report["day"].opened_at.strftime("%d/%m/%Y %H:%M:%S") if report["day"].opened_at else "—"])
    ws.append(["Cierre del día", report["day"].closed_at.strftime("%d/%m/%Y %H:%M:%S") if report["day"].closed_at else "Pendiente"])
    ws.append([])
    ws.append(["Concepto", "Valor"])
    ws.append(["Total esperado cajas", float(report["total_expected"])])
    ws.append(["Total contado cajas", float(report["total_counted"])])
    ws.append(["Total diferencia", float(report["total_difference"])])
    ws.append(["Efectivo ventas", float(report["cash_sales"])])
    ws.append(["Efectivo pagos de cartera", float(report["cash_payments"])])
    ws.append(["Reversiones efectivo ventas", float(report["cash_sales_reversals"])])
    ws.append(["Reversiones efectivo pagos", float(report["cash_payment_reversals"])])
    ws.append([])
    ws.append(["VENTAS DEL DÍA POR MEDIO DE PAGO"])
    ws.append(["Medio", "Total"])
    for method in report["methods"]:
        ws.append([method, float(report["sales_day"].get(method, ZERO))])
    wp = wb.create_sheet("Pagos del día")
    wp.append(["MEDIO", "TOTAL PAGOS"])
    for method in report["methods"]:
        wp.append([method, float(report["payments_day"].get(method, ZERO))])
    wc = wb.create_sheet("Ventas por caja")
    wc.append(["Sucursal", "Caja", "Estado", *report["methods"], "Total ventas"])
    for row in report["registers"]:
        values = [row["sales"].get(method, ZERO) for method in report["methods"]]
        wc.append([row["branch"], row["box"], row["status"], *map(float, values), float(sum(values, ZERO))])
    wpc = wb.create_sheet("Pagos por caja")
    wpc.append(["Sucursal", "Caja", "Estado", *report["methods"], "Total pagos"])
    for row in report["registers"]:
        values = [row["payments"].get(method, ZERO) for method in report["methods"]]
        wpc.append([row["branch"], row["box"], row["status"], *map(float, values), float(sum(values, ZERO))])
    wa = wb.create_sheet("Arqueos")
    wa.append(["Sucursal", "Caja", "Estado", "Base", "Apertura", "Hora apertura", "Hora cierre", "Esperado", "Contado", "Diferencia", "Resultado"])
    for row in report["registers"]:
        result = "OK" if row["difference"] == ZERO else "DESCUADRADA" if row["difference"] is not None else "PENDIENTE"
        wa.append([
            row["branch"], row["box"], row["status"], float(row["base_amount"]), float(row["opening_amount"]),
            row["opened_at"].strftime("%d/%m/%Y %H:%M:%S") if row["opened_at"] else "—",
            row["closed_at"].strftime("%d/%m/%Y %H:%M:%S") if row["closed_at"] else "Pendiente",
            float(row["expected"]), None if row["counted"] is None else float(row["counted"]),
            None if row["difference"] is None else float(row["difference"]), result,
        ])
    wm = wb.create_sheet("Movimientos efectivo")
    wm.append(["Concepto", "Total"])
    wm.append(["Efectivo a caja por ventas", float(report["cash_sales"])])
    wm.append(["Efectivo a caja por pagos", float(report["cash_payments"])])
    wm.append(["Reversiones efectivo ventas", float(report["cash_sales_reversals"])])
    wm.append(["Reversiones efectivo pagos", float(report["cash_payment_reversals"])])
    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
        for column in sheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column) + 2, 35)
            sheet.column_dimensions[column[0].column_letter].width = width
    output = BytesIO()
    wb.save(output)
    return output.getvalue(), _filename(report, "xlsx")


def _pdf_table(rows):
    table = Table(rows, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9EAF7")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
        ("ALIGN", (2, 1), (-1, -1), "RIGHT"),
    ]))
    return table


def pdf_report(report) -> tuple[bytes, str]:
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(A4), rightMargin=10 * mm, leftMargin=10 * mm, topMargin=10 * mm, bottomMargin=10 * mm)
    styles = getSampleStyleSheet()
    day = report["day"]
    generated_at = report["generated_at"].astimezone(COLOMBIA_TZ) if report["generated_at"].tzinfo else report["generated_at"]
    story = [
        Paragraph("Resumen de cierre de Caja", styles["Title"]),
        Paragraph(f"Tenant: {report['tenant_name']}", styles["Normal"]),
        Paragraph(f"Fecha operativa: {day.business_date} · Estado: {day.status}", styles["Normal"]),
        Paragraph(f"Generado por: {report['generated_by']} · Generado: {generated_at.strftime('%d/%m/%Y %H:%M:%S')}", styles["Normal"]),
        Paragraph(f"Apertura del día: {day.opened_at.strftime('%d/%m/%Y %H:%M:%S') if day.opened_at else '—'} · Cierre del día: {day.closed_at.strftime('%d/%m/%Y %H:%M:%S') if day.closed_at else 'Pendiente'}", styles["Normal"]),
        Spacer(1, 5 * mm),
    ]
    summary = [
        ["Total esperado", "Total contado", "Total diferencia", "Efectivo ventas", "Efectivo pagos"],
        [f"${report['total_expected']:,.0f}", f"${report['total_counted']:,.0f}", f"${report['total_difference']:,.0f}", f"${report['cash_sales']:,.0f}", f"${report['cash_payments']:,.0f}"],
    ]
    table = Table(summary, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9EAF7")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
    ]))
    story.extend([table, Spacer(1, 5 * mm)])
    methods = report["methods"]
    sales_table = [["Caja", "Estado", *methods, "Total"]]
    for row in report["registers"]:
        values = [row["sales"].get(method, ZERO) for method in methods]
        sales_table.append([row["box"], row["status"], *[f"${value:,.0f}" for value in values], f"${sum(values, ZERO):,.0f}"])
    story.extend([Paragraph("Ventas por caja", styles["Heading2"]), _pdf_table(sales_table), Spacer(1, 5 * mm)])
    payment_table = [["Caja", "Estado", *methods, "Total"]]
    for row in report["registers"]:
        values = [row["payments"].get(method, ZERO) for method in methods]
        payment_table.append([row["box"], row["status"], *[f"${value:,.0f}" for value in values], f"${sum(values, ZERO):,.0f}"])
    story.extend([Paragraph("Pagos de cartera por caja", styles["Heading2"]), _pdf_table(payment_table), Spacer(1, 5 * mm)])
    reconciliation = [["Caja", "Sucursal", "Base", "Apertura", "Hora apertura", "Hora cierre", "Esperado", "Contado", "Diferencia", "Resultado"]]
    for row in report["registers"]:
        result = "OK" if row["difference"] == ZERO else "DESCUADRADA" if row["difference"] is not None else "PENDIENTE"
        reconciliation.append([
            row["box"], row["branch"], f"${row['base_amount']:,.0f}", f"${row['opening_amount']:,.0f}",
            row["opened_at"].strftime("%d/%m/%Y %H:%M:%S") if row["opened_at"] else "—",
            row["closed_at"].strftime("%d/%m/%Y %H:%M:%S") if row["closed_at"] else "Pendiente",
            f"${row['expected']:,.0f}", "—" if row["counted"] is None else f"${row['counted']:,.0f}",
            "—" if row["difference"] is None else f"${row['difference']:,.0f}", result,
        ])
    story.extend([Paragraph("Arqueos y diferencias", styles["Heading2"]), _pdf_table(reconciliation), Spacer(1, 4 * mm)])
    if report["pending"]:
        story.append(Paragraph(f"Cajas pendientes de arqueo: {len(report['pending'])}", styles["Normal"]))
    credit_sales = report["sales_day"].get("Crédito", ZERO)
    if credit_sales:
        story.append(Paragraph("Ventas a crédito del día (sin movimiento de caja asignable): " f"${credit_sales:,.0f}", styles["Normal"]))
    doc.build(story)
    return output.getvalue(), _filename(report, "pdf")
